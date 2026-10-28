"""Retirada autenticada e recebimento independente, apenas em banco descartável."""
import asyncio
from datetime import datetime, timedelta, timezone
from io import BytesIO
import os
import secrets
import unittest
from unittest.mock import patch
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import func, select, update
from sqlalchemy.ext.asyncio import create_async_engine

import test_entregas as fixtures
from app.models.models import EventoAssistencial, Pedido, RecebimentoQR, RetiradaPedido, SessaoCopeira, Usuario


class RetiradaEntregaTest(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.f = fixtures.EntregasTest('test_dashboard_abre_sem_login_e_publica_apenas_dados_operacionais')
        await self.f.asyncSetUp()
        self.id = self.f.pedidos['PRONTO']
        self.base = f'/api/operacao/unidades/{self.f.a}/pedidos/{self.id}'
        self.codigo = secrets.token_urlsafe(32)
        async with self.f.sessoes() as db:
            db.add(RecebimentoQR(pedido_id=self.id, versao=1, codigo=self.codigo, url_base='https://teste.invalid'))
            await db.commit()
        self.ana = await self.f.entrar('ana')
        self.bia = await self.f.entrar('bia')
        self.admin = await self.f.entrar('admin')

    async def asyncTearDown(self):
        await self.f.asyncTearDown()

    async def retirar(self, usuario=None, dados=None):
        return await self.f.client.post(self.base + '/retirar', headers=usuario or self.ana, json=dados or {})

    async def confirmar(self):
        return await self.f.client.post('/api/recebimentos/confirmar', headers=self.ana, json={'codigo': self.codigo,
            'recebedor_nome': 'Recebedora Fictícia', 'registro_profissional': 'COREN-SP 000001'})

    async def test_primeira_retirada_exige_login_identidade_propria_unidade_e_presenca(self):
        self.assertEqual((await self.f.client.post(self.base + '/retirar', json={})).status_code, 401)
        self.assertEqual((await self.retirar(dados={'copeira_id': self.f.bia})).status_code, 403)
        gestor = await self.f.entrar('gestor')
        self.assertEqual((await self.retirar(gestor, {'copeira_id': self.f.ana})).status_code, 403)
        for dados in ({'retirada_em': '2000-01-01'}, {'copeira_nome': 'Outra pessoa'}):
            self.assertEqual((await self.retirar(dados=dados)).status_code, 422)
        async with self.f.sessoes() as db:
            await db.execute(update(SessaoCopeira).where(SessaoCopeira.usuario_id == self.f.ana)
                .values(presenca_ate=datetime.now(timezone.utc).replace(tzinfo=None) - timedelta(minutes=1)))
            await db.commit()
        self.assertEqual((await self.retirar()).status_code, 409)

    async def test_retirada_rejeita_versao_desatualizada_pedido_de_outra_unidade_e_etapa_invalida(self):
        self.assertEqual((await self.retirar(dados={'versao': 2})).status_code, 409)
        outra = f'/api/operacao/unidades/{self.f.b}/pedidos/{self.id}/retirar'
        self.assertEqual((await self.f.client.post(outra, headers=self.ana, json={})).status_code, 404)
        for etapa in ('PENDENTE', 'EM_PREPARO', 'CANCELADO'):
            url = f'/api/operacao/unidades/{self.f.a}/pedidos/{self.f.pedidos[etapa]}/retirar'
            self.assertEqual((await self.f.client.post(url, headers=self.ana, json={})).status_code, 409)
        self.assertEqual((await self.retirar(dados={'versao': 1})).status_code, 201)

    async def test_retirada_nao_entrega_reenvio_nao_duplica_e_outra_copeira_nao_assume(self):
        primeira = await self.retirar()
        self.assertEqual(primeira.status_code, 201, primeira.text)
        self.assertEqual(primeira.json()['status'], 'PRONTO')
        retirada = primeira.json()['retirada']
        self.assertEqual(retirada['copeira_id'], self.f.ana)
        repetida = await self.retirar()
        self.assertEqual(repetida.status_code, 200, repetida.text)
        self.assertEqual(repetida.json()['retirada'], retirada)
        self.assertEqual((await self.retirar(self.bia)).status_code, 409)
        async with self.f.sessoes() as db:
            self.assertEqual(await db.scalar(select(func.count()).select_from(RetiradaPedido)), 1)
            self.assertEqual(await db.scalar(select(func.count()).select_from(EventoAssistencial).where(EventoAssistencial.acao == 'RETIRADA_REGISTRADA')), 1)
            self.assertIsNone((await db.get(RecebimentoQR, self.id)).confirmado_em)
        publica = await self.f.client.get(f'/api/operacao/unidades/{self.f.a}/estado')
        item = next(p for p in publica.json()['pedidos'] if p['id'] == self.id)
        self.assertTrue(item['em_entrega'])
        self.assertNotIn('Ana Teste', publica.text)
        self.assertNotIn('retirada_em', publica.text)

    async def test_qr_preserva_identificacao_apos_logout_e_exige_nova_sessao_para_confirmar(self):
        antes = await self.f.client.post('/api/recebimentos/consultar', json={'codigo': self.codigo})
        self.assertFalse(antes.json()['pode_confirmar'])
        self.assertTrue(antes.json()['aguardando_retirada'])
        self.assertEqual((await self.confirmar()).status_code, 409)
        await self.retirar()
        await self.f.client.post('/api/auth/logout', headers=self.ana)
        async with self.f.sessoes() as db:
            (await db.get(Usuario, self.f.ana)).nome = 'Nome atualizado depois'
            await db.commit()
        depois = await self.f.client.post('/api/recebimentos/consultar', json={'codigo': self.codigo})
        self.assertEqual(depois.json()['copeira_nome'], 'Ana Teste')
        self.assertTrue(depois.json()['pode_confirmar'])
        self.assertEqual((await self.confirmar()).status_code, 401)
        self.ana = await self.f.entrar('ana')
        primeira = await self.confirmar()
        self.assertEqual(primeira.status_code, 200, primeira.text)
        self.assertEqual(primeira.json()['status'], 'ENTREGUE')
        self.assertEqual(primeira.json(), (await self.confirmar()).json())
        async with self.f.sessoes() as db:
            qr = await db.get(RecebimentoQR, self.id)
            self.assertEqual(qr.recebedor_nome, 'Recebedora Fictícia')
            self.assertGreaterEqual(qr.confirmado_em, (await db.get(RetiradaPedido, (self.id, 1))).retirada_em)

    async def test_gestao_indica_copeira_e_relatorios_exportam_dois_horarios(self):
        self.assertEqual((await self.retirar(self.admin)).status_code, 422)
        self.assertEqual((await self.retirar(self.admin, {'copeira_id': self.f.ana})).status_code, 201)
        filtro = {'inicio': '2020-01-01T00:00:00Z', 'fim': '2100-01-01T00:00:00Z'}
        r = await self.f.client.get('/api/relatorios/dietas-periodo', params=filtro, headers=self.admin)
        item = next(p for p in r.json()['itens'] if p['pedido_id'] == self.id)
        self.assertEqual(item['etapa_entrega'], 'EM_ENTREGA')
        self.assertIsNone(item['entregue_em'])
        self.assertIsNotNone(item['retirada_em'])
        await self.confirmar()
        for formato in ('xlsx', 'ods', 'csv'):
            exportacao = await self.f.client.get('/api/relatorios/dietas-periodo/exportar', params={**filtro, 'formato': formato}, headers=self.admin)
            self.assertEqual(exportacao.status_code, 200, exportacao.text[:100] if formato == 'csv' else '')
            if formato == 'csv':
                self.assertIn('Identificação da copeira (Brasília)', exportacao.text)
                self.assertIn('Ana Teste', exportacao.text)
                self.assertIn('Recebedora Fictícia', exportacao.text)
            if formato == 'xlsx':
                livro = load_workbook(BytesIO(exportacao.content))
                folha = livro['Solicitações']
                colunas = {c.value: c.column for c in folha[7]}
                linha = next(n for n in range(8, folha.max_row + 1) if folha.cell(n, colunas['Pedido']).value == self.id)
                self.assertIsInstance(folha.cell(linha, colunas['Identificação da copeira (Brasília)']).value, datetime)
                self.assertIsInstance(folha.cell(linha, colunas['Entrega (Brasília)']).value, datetime)
        entregas = await self.f.client.get('/api/relatorios/entregas', params={'filial_id': self.f.a}, headers=self.admin)
        item = next(p for p in entregas.json()['itens'] if p['pedido_id'] == self.id)
        self.assertEqual(item['copeira_nome'], 'Ana Teste')
        self.assertIsNotNone(item['retirada_em'])

    async def test_versao_corrigida_nao_herda_copeira_de_filipeta_anterior(self):
        await self.retirar()
        async with self.f.sessoes() as db:
            (await db.get(Pedido, self.id)).versao = 2
            (await db.get(RecebimentoQR, self.id)).versao = 2
            await db.commit()
        consulta = await self.f.client.post('/api/recebimentos/consultar', json={'codigo': self.codigo})
        self.assertIsNone(consulta.json()['copeira_nome'])
        self.assertEqual((await self.confirmar()).status_code, 409)
        nova = await self.retirar(self.bia)
        self.assertEqual(nova.status_code, 201, nova.text)
        self.assertEqual(nova.json()['retirada']['copeira_id'], self.f.bia)
        async with self.f.sessoes() as db:
            self.assertEqual(await db.scalar(select(func.count()).select_from(RetiradaPedido)), 2)

    async def identificar(self, cabecalho=None):
        return await self.f.client.post('/api/recebimentos/identificar', headers=cabecalho or self.ana, json={'codigo': self.codigo})

    async def test_login_e_identificacao_acontecem_no_qr_sem_etapa_no_painel(self):
        sem_login = await self.f.client.post('/api/recebimentos/identificar', json={'codigo': self.codigo})
        self.assertEqual(sem_login.status_code, 401)
        primeira = await self.identificar()
        self.assertEqual(primeira.status_code, 200, primeira.text)
        self.assertTrue(primeira.json()['pode_confirmar'])
        self.assertEqual(primeira.json()['copeira_nome'], 'Ana Teste')
        self.assertEqual(primeira.json(), (await self.identificar()).json())
        self.assertEqual((await self.identificar(self.bia)).status_code, 409)
        self.assertEqual((await self.f.client.post('/api/recebimentos/confirmar', headers=self.bia,
            json={'codigo': self.codigo, 'recebedor_nome': 'Pessoa Fictícia', 'registro_profissional': 'COREN 123'})).status_code, 403)
        self.assertEqual((await self.confirmar()).status_code, 200)
        async with self.f.sessoes() as db:
            self.assertEqual(await db.scalar(select(func.count()).select_from(RetiradaPedido)), 1)
            self.assertEqual(await db.scalar(select(func.count()).select_from(EventoAssistencial).where(EventoAssistencial.acao == 'COPEIRA_IDENTIFICADA_QR')), 1)

    async def test_qr_exige_copeira_da_unidade_e_rejeita_identidade_ou_horario_enviados_pelo_cliente(self):
        self.assertEqual((await self.identificar(self.admin)).status_code, 403)
        for dados in ({'copeira_id': self.f.bia}, {'identificada_em': '2000-01-01'}, {'copeira_nome': 'Outra pessoa'}):
            r = await self.f.client.post('/api/recebimentos/identificar', headers=self.ana, json={'codigo': self.codigo, **dados})
            self.assertEqual(r.status_code, 422)
        async with self.f.sessoes() as db:
            (await db.get(Pedido, self.id)).filial_id = self.f.b
            await db.commit()
        self.assertEqual((await self.identificar(self.bia)).status_code, 403)


@unittest.skipUnless(os.environ.get('CNUTRI_TEST_POSTGRES_URL'), 'Requer PostgreSQL de teste')
class RetiradaEntregaPostgresTest(RetiradaEntregaTest):
    async def asyncSetUp(self):
        url = os.environ['CNUTRI_TEST_POSTGRES_URL']
        self.schema = 'teste_retirada_' + uuid4().hex
        self.admin_engine = create_async_engine(url)
        async with self.admin_engine.begin() as conn:
            await conn.exec_driver_sql(f'CREATE SCHEMA "{self.schema}"')
        def motor_descartavel(*_args, **_kwargs):
            return create_async_engine(url, connect_args={'server_settings': {'search_path': self.schema}})
        with patch.object(fixtures, 'create_async_engine', motor_descartavel):
            await super().asyncSetUp()

    async def asyncTearDown(self):
        await super().asyncTearDown()
        async with self.admin_engine.begin() as conn:
            await conn.exec_driver_sql(f'DROP SCHEMA "{self.schema}" CASCADE')
        await self.admin_engine.dispose()

    async def test_duas_copeiras_simultaneas_nao_trocam_responsavel(self):
        resultados = await asyncio.gather(self.retirar(self.ana), self.retirar(self.bia))
        self.assertEqual(sorted(r.status_code for r in resultados), [201, 409], [r.text for r in resultados])
        async with self.f.sessoes() as db:
            self.assertEqual(await db.scalar(select(func.count()).select_from(RetiradaPedido)), 1)
            self.assertEqual(await db.scalar(select(func.count()).select_from(EventoAssistencial).where(EventoAssistencial.acao == 'RETIRADA_REGISTRADA')), 1)

    async def test_reenvios_simultaneos_da_mesma_copeira_confirmam_um_unico_registro(self):
        resultados = await asyncio.gather(self.retirar(), self.retirar())
        self.assertEqual(sorted(r.status_code for r in resultados), [200, 201], [r.text for r in resultados])
        self.assertEqual(resultados[0].json()['retirada'], resultados[1].json()['retirada'])

    async def test_duas_leituras_qr_de_copeiras_distintas_identificam_uma_so(self):
        resultados = await asyncio.gather(self.identificar(self.ana), self.identificar(self.bia))
        self.assertEqual(sorted(r.status_code for r in resultados), [200, 409], [r.text for r in resultados])
        async with self.f.sessoes() as db:
            self.assertEqual(await db.scalar(select(func.count()).select_from(RetiradaPedido)), 1)
            self.assertEqual(await db.scalar(select(func.count()).select_from(EventoAssistencial).where(EventoAssistencial.acao == 'COPEIRA_IDENTIFICADA_QR')), 1)
