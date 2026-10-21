"""Relatórios/exportações com dados fictícios e banco isolado, sem acessar pacientes reais."""
import csv
from datetime import datetime, timedelta, timezone
from uuid import uuid4
from io import BytesIO, StringIO
import unittest

import httpx
from odf.opendocument import load
from odf.table import Table, TableCell, TableRow
from odf.teletype import extractText
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.core.database import Base, get_db
from app.core.security import criar_token, hash_senha
from app.main import app
from app.models.models import Dieta, EntregaPedido, Filial, Pedido, Perfil, RecebimentoEntrega, SessaoUsuario, Usuario
from app.services.exportacao_dietas import COLUNAS

HASH = hash_senha("TesteSeguro@123")


class RelatorioDietasTest(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.engine = create_async_engine("sqlite+aiosqlite:///:memory:")
        self.sessoes = async_sessionmaker(self.engine, expire_on_commit=False)
        async with self.engine.begin() as conexao:
            await conexao.run_sync(Base.metadata.create_all)
        self.ids = {}
        async with self.sessoes() as db:
            a = Filial(nome="Hospital A · Teste", codigo_slug="hospital-a")
            b = Filial(nome="Hospital B · Teste", codigo_slug="hospital-b")
            c = Filial(nome="Hospital C · Inativo", codigo_slug="hospital-c", ativa=False)
            perfis = {nome: Perfil(nome=nome) for nome in ("ADMIN_TI", "NUTRICIONISTA", "COZINHA")}
            db.add_all([a, b, c, *perfis.values()])
            await db.flush()
            self.a, self.b, self.c = a.id, b.id, c.id
            geral = Dieta(nome="GERAL", filial_id=a.id)
            pastosa = Dieta(nome="PASTOSA · Antiga", filial_id=b.id, ativa=False)
            usuarios = {
                "admin": Usuario(nome="Admin Teste", nome_usuario="admin", email="admin@teste.local", senha_hash=HASH, perfil=perfis["ADMIN_TI"], filiais=[]),
                "gestor": Usuario(nome="Gestor A", nome_usuario="gestor", email="gestor@teste.local", senha_hash=HASH, perfil=perfis["NUTRICIONISTA"], filiais=[a]),
                "cozinha": Usuario(nome="Cozinha Teste", nome_usuario="cozinha", email="cozinha@teste.local", senha_hash=HASH, perfil=perfis["COZINHA"], filiais=[a]),
            }
            db.add_all([geral, pastosa, *usuarios.values()])
            await db.flush()
            self.tokens = {}
            for nome, usuario in usuarios.items():
                jti = str(uuid4())
                db.add(SessaoUsuario(id=jti, usuario_id=usuario.id, expira_em=datetime.now(timezone.utc).replace(tzinfo=None) + timedelta(hours=1)))
                self.tokens[nome] = {"Authorization": "Bearer " + criar_token(usuario, jti)}

            async def pedido(chave, unidade, horario, status="PENDENTE", dieta=None, medico=False, nome="João · Paciente fictício"):
                item = Pedido(filial_id=unidade.id, criado_em=horario, status=status, dieta_id=dieta.id if dieta else None,
                    tipo_solicitacao="MEDICO" if medico else "PACIENTE", solicitante_nome="Solicitante Fictício",
                    solicitante_identificacao="PRIVADO", nome_paciente=None if medico else nome, leito="0012",
                    dados_dinamicos={"refeicao": "ALMOÇO, opção com açafrão", "cardapio": "+TEXTO DE TESTE",
                        "suco": "COM SUCO", "andar": "TÉRREO", "matricula": "PRIVADO", "data_nascimento": "1970-01-01"},
                    observacoes_alergias="ação, maçã e coração\nSEM SAL")
                db.add(item)
                await db.flush()
                self.ids[chave] = item.id
                return item

            await pedido("antes", a, datetime(2026, 10, 20, 14, 59, 59), dieta=geral)
            await pedido("inicio", a, datetime(2026, 10, 20, 15), dieta=geral)
            await pedido("pronto", a, datetime(2026, 10, 20, 15, 30), "PRONTO", geral)
            await pedido("fim", a, datetime(2026, 10, 20, 16), "CANCELADO", geral)
            await pedido("depois", a, datetime(2026, 10, 20, 16, 0, 0, 1), dieta=geral)
            medico = await pedido("medico", a, datetime(2026, 10, 20, 15, 45), "ENTREGUE", medico=True)
            await pedido("outro", b, datetime(2026, 10, 20, 15, 10), "EM_PREPARO", pastosa)
            await pedido("inativo", c, datetime(2026, 10, 20, 15, 20), medico=True)
            for n in range(26):
                await pedido(f"extra{n}", a, datetime(2026, 10, 20, 15, 5) + timedelta(seconds=n), dieta=geral,
                    nome='=HYPERLINK("https://example.invalid","TESTE")' if n == 0 else "Paciente Fictício")
            db.add(EntregaPedido(pedido_id=medico.id, copeira_id=usuarios["gestor"].id, copeira_nome="Copeira Fictícia", entregue_em=datetime(2026, 10, 20, 17, 30)))
            await db.flush()
            db.add(RecebimentoEntrega(pedido_id=medico.id, recebedor_nome="Recebedor Fictício", assinatura_png=b"assinatura de teste"))
            await db.commit()

        async def banco_teste():
            async with self.sessoes() as db:
                yield db
        app.dependency_overrides[get_db] = banco_teste
        self.client = httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://teste")
        self.base = "/api/relatorios/dietas-periodo"
        self.periodo = {"inicio": "2026-10-20T12:00:00-03:00", "fim": "2026-10-20T13:00:00-03:00"}

    async def asyncTearDown(self):
        await self.client.aclose()
        app.dependency_overrides.clear()
        await self.engine.dispose()

    async def obter(self, *, token="admin", **filtros):
        resposta = await self.client.get(self.base, params={**self.periodo, **filtros}, headers=self.tokens[token])
        self.assertEqual(resposta.status_code, 200, resposta.text)
        return resposta.json()

    async def arquivo(self, formato, **filtros):
        return await self.client.get(self.base + "/exportar", params={**self.periodo, "formato": formato, **filtros}, headers=self.tokens["admin"])

    async def test_resumo_periodo_inclusivo_fuso_e_paginacao(self):
        resultado = await self.obter()
        self.assertEqual(resultado["totais"], {"total": 32, "pendentes": 28, "em_preparo": 1, "prontos": 1, "entregues": 1, "cancelados": 1})
        self.assertEqual(sum(r["total"] for r in resultado["resumo"]), 32)
        self.assertEqual(len(resultado["itens"]), 25)
        pagina2 = await self.obter(pagina=2)
        self.assertEqual(len(pagina2["itens"]), 7)
        todos = resultado["itens"] + pagina2["itens"]
        ids = {p["pedido_id"] for p in todos}
        self.assertIn(self.ids["inicio"], ids)
        self.assertIn(self.ids["fim"], ids)
        self.assertNotIn(self.ids["antes"], ids)
        self.assertNotIn(self.ids["depois"], ids)
        self.assertEqual((await self.obter(inicio="2026-10-20T12:00:00", fim="2026-10-20T13:00:00"))["totais"], resultado["totais"])

    async def test_permissoes_da_consulta_e_dos_tres_formatos(self):
        self.assertEqual((await self.obter(token="gestor"))["totais"]["total"], 30)
        self.assertEqual((await self.obter(filial_id=self.b))["totais"]["total"], 1)
        self.assertTrue(any(r["dieta"] == "PASTOSA · Antiga" for r in (await self.obter())["resumo"]))
        self.assertEqual((await self.obter(filial_id=self.c))["totais"]["total"], 1)
        for formato in (None, "xlsx", "ods", "csv"):
            rota = self.base + ("/exportar" if formato else "")
            parametros = {**self.periodo, **({"formato": formato} if formato else {})}
            self.assertEqual((await self.client.get(rota, params=parametros)).status_code, 401)
            self.assertEqual((await self.client.get(rota, params=parametros, headers=self.tokens["cozinha"])).status_code, 403)
            self.assertEqual((await self.client.get(rota, params={**parametros, "filial_id": self.b}, headers=self.tokens["gestor"])).status_code, 403)

    async def test_filtra_formulario_status_e_data_real_da_entrega(self):
        self.assertEqual((await self.obter(tipo_solicitacao="MEDICO"))["totais"]["total"], 2)
        resultado = await self.obter(status="ENTREGUE", tipo_solicitacao="MEDICO", data_base="entrega",
            inicio="2026-10-20T14:29:00-03:00", fim="2026-10-20T14:31:00-03:00")
        self.assertEqual(resultado["totais"]["total"], 1)
        self.assertEqual(resultado["itens"][0]["pedido_id"], self.ids["medico"])
        self.assertEqual(resultado["itens"][0]["recebedor_nome"], "Recebedor Fictício")
        self.assertTrue(resultado["itens"][0]["assinatura_registrada"])
        self.assertEqual((await self.obter(data_base="entrega"))["totais"]["total"], 0)

    async def test_xlsx_inclui_todos_os_registros_datas_e_textos_sem_formulas(self):
        arquivo = await self.arquivo("xlsx")
        self.assertEqual(arquivo.status_code, 200, arquivo.text[:200] if arquivo.status_code != 200 else "")
        livro = load_workbook(BytesIO(arquivo.content), read_only=False)
        self.assertEqual(livro.sheetnames, ["Resumo", "Solicitações"])
        folha = livro["Solicitações"]
        self.assertEqual(folha.max_row, 39)
        self.assertEqual(folha.freeze_panes, "C8")
        indices = {chave: n for n, (chave, _, _) in enumerate(COLUNAS, 1)}
        celulas = list(folha.iter_rows(min_row=8))
        inicio = next(l for l in celulas if l[0].value == self.ids["inicio"])
        self.assertEqual(inicio[indices["criado_em"] - 1].value, datetime(2026, 10, 20, 12))
        self.assertEqual(inicio[indices["leito"] - 1].value, "0012")
        self.assertIn("coração\nSEM SAL", inicio[indices["observacoes_alergias"] - 1].value)
        perigoso = next(l for l in celulas if l[0].value == self.ids["extra0"])
        self.assertEqual(perigoso[indices["destinatario"] - 1].data_type, "s")
        self.assertTrue(perigoso[indices["destinatario"] - 1].value.startswith("=HYPERLINK"))
        self.assertFalse(any(c.data_type == "f" for linha in celulas for c in linha))
        resumo = list(livro["Resumo"].iter_rows(min_row=8, values_only=True))
        self.assertEqual(resumo[-1][2], 32)
        livro.close()

    async def test_ods_e_csv_abrem_com_acentos_e_sem_formulas(self):
        ods = await self.arquivo("ods")
        self.assertEqual(ods.status_code, 200, ods.text[:200] if ods.status_code != 200 else "")
        livro = load(BytesIO(ods.content))
        folhas = livro.spreadsheet.getElementsByType(Table)
        self.assertEqual([t.getAttribute("name") for t in folhas], ["Resumo", "Solicitações"])
        linhas = folhas[1].getElementsByType(TableRow)[7:]
        self.assertEqual(len(linhas), 32)
        self.assertFalse(any(c.getAttribute("formula") for linha in linhas for c in linha.getElementsByType(TableCell)))
        csv_arquivo = await self.arquivo("csv")
        self.assertEqual(csv_arquivo.status_code, 200)
        self.assertTrue(csv_arquivo.content.startswith(b"\xef\xbb\xbf"))
        csv_linhas = list(csv.DictReader(StringIO(csv_arquivo.content.decode("utf-8-sig"))))
        self.assertEqual(len(csv_linhas), 32)
        self.assertTrue(any(l["Destinatário"].startswith("'=HYPERLINK") for l in csv_linhas))
        self.assertTrue(all(l["Cardápio"].startswith("'+") for l in csv_linhas))
        self.assertTrue(all("ação, maçã e coração\nSEM SAL" == l["Restrições / observações"] for l in csv_linhas))
        self.assertIn("no-store", csv_arquivo.headers["cache-control"])
        self.assertIn(".csv", csv_arquivo.headers["content-disposition"])

    async def test_filtros_invalidos_e_relatorio_vazio(self):
        for filtros in ({"inicio": "2026-10-21T00:00:00-03:00"}, {"pagina": 0}, {"data_base": "outra"}, {"status": "INVALIDO"}):
            resposta = await self.client.get(self.base, params={**self.periodo, **filtros}, headers=self.tokens["admin"])
            self.assertEqual(resposta.status_code, 422, resposta.text)
        vazio = await self.obter(inicio="2026-10-19T00:00:00-03:00", fim="2026-10-19T23:59:59-03:00")
        self.assertEqual(vazio["totais"]["total"], 0)
        self.assertEqual(vazio["itens"], [])
        for formato in ("xlsx", "ods", "csv"):
            self.assertEqual((await self.arquivo(formato, inicio="2026-10-19T00:00:00-03:00", fim="2026-10-19T23:59:59-03:00")).status_code, 200)

    async def test_texto_muito_longo_nao_e_truncado_silenciosamente(self):
        async with self.sessoes() as db:
            (await db.get(Pedido, self.ids["inicio"])).observacoes_alergias = "á" * 33000
            await db.commit()
        xlsx = await self.arquivo("xlsx")
        self.assertEqual(xlsx.status_code, 400)
        self.assertIn("CSV ou ODS", xlsx.text)
        csv_arquivo = await self.arquivo("csv")
        linhas = list(csv.DictReader(StringIO(csv_arquivo.content.decode("utf-8-sig"))))
        self.assertTrue(any(len(l["Restrições / observações"]) == 33000 for l in linhas))


if __name__ == "__main__":
    unittest.main()
