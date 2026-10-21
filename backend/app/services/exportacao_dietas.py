"""Exportadores da aplicação: XLSX, ODS e CSV sem fórmulas vindas dos formulários."""
import csv
from datetime import datetime
from io import BytesIO, StringIO
from math import ceil
import re

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from odf.opendocument import OpenDocumentSpreadsheet
from odf.number import DateStyle, Day, Hours, Minutes, Month, Seconds, Text, Year
from odf.style import Style, TableCellProperties, TableColumnProperties, TableRowProperties, TextProperties
from odf.table import Table, TableCell, TableColumn, TableRow
from odf.text import P

from .relatorio_dietas import STATUS, TIPOS, horario_brasilia

COLUNAS = [
    ("pedido_id", "Pedido", 11), ("unidade", "Unidade", 38), ("criado_em", "Solicitação (Brasília)", 24),
    ("tipo_solicitacao", "Formulário", 16), ("status", "Situação", 16), ("destinatario", "Destinatário", 32),
    ("leito", "Leito", 14), ("local_entrega", "Local de entrega", 32), ("dieta", "Dieta", 30),
    ("refeicao", "Refeição", 48), ("cardapio", "Cardápio", 48), ("para", "Para", 28),
    ("equipe", "Equipe", 20), ("horario_entrega", "Horário previsto", 18), ("suco", "Suco", 20),
    ("sobremesa", "Sobremesa", 26), ("liquido_espessado", "Líquido espessado", 22),
    ("internado_mais_6h", "Internado há mais de 6h", 24), ("observacoes_alergias", "Restrições / observações", 52),
    ("solicitante_nome", "Solicitante", 32), ("concluido_em", "Preparo concluído (Brasília)", 26),
    ("entregue_em", "Entrega (Brasília)", 24), ("copeira_nome", "Copeira", 30), ("recebedor_nome", "Recebedor", 30),
    ("assinatura_registrada", "Assinatura registrada", 23), ("outros_campos", "Outros campos informados", 52),
    ("registrado_por_nome", "Entrega registrada por", 32),
]
RESUMO = [("unidade", "Unidade", 38), ("dieta", "Dieta / categoria", 32), ("total", "Solicitações", 15),
    ("pendentes", "Pendentes", 14), ("em_preparo", "Em preparo", 14), ("prontos", "Prontos", 13),
    ("entregues", "Entregues", 13), ("cancelados", "Cancelados", 14)]
CONTROLES = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
MIMES = {"xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "ods": "application/vnd.oasis.opendocument.spreadsheet", "csv": "text/csv"}


def limpar(valor):
    return CONTROLES.sub("", valor) if isinstance(valor, str) else valor


def valor_celula(item, chave):
    valor = item.get(chave)
    if chave in ("criado_em", "concluido_em", "entregue_em"):
        return horario_brasilia(valor)
    if chave == "status":
        return STATUS.get(valor, valor)
    if chave == "tipo_solicitacao":
        return TIPOS.get(valor, valor)
    if isinstance(valor, bool):
        return "Sim" if valor else "Não"
    return limpar(valor)


def tabelas(itens, resumo, totais):
    total = {"unidade": "TOTAL", "dieta": "", **totais}
    return [("Resumo", RESUMO, [*resumo, total]), ("Solicitações", COLUNAS, itens)]


def contexto_linhas(contexto):
    inicio = datetime.fromisoformat(contexto["inicio"]).strftime("%d/%m/%Y %H:%M:%S")
    fim = datetime.fromisoformat(contexto["fim"]).strftime("%d/%m/%Y %H:%M:%S")
    criterio = "solicitação" if contexto["data_base"] == "solicitacao" else "entrega"
    return ["Cnutri — Relatório de dietas e solicitações", f"Período por {criterio}: {inicio} a {fim} (Brasília)",
        "Unidades: " + (", ".join(contexto["unidades"]) or "Nenhuma unidade disponível"),
        f'Formulário: {contexto["tipo"]}. Situação: {contexto["status"]}. Uma linha por solicitação.',
        "Fonte: registros do Cnutri. Os horários das planilhas estão em Brasília (UTC-3)."]


def xlsx(itens, resumo, totais, contexto):
    livro = Workbook(write_only=True)
    for nome, colunas, registros in tabelas(itens, resumo, totais):
        folha = livro.create_sheet(nome)
        folha.sheet_view.showGridLines = False
        folha.freeze_panes = "C8"
        folha.auto_filter.ref = f"A7:{get_column_letter(len(colunas))}{max(7, 7 + len(registros))}"
        for n, (_, _, largura) in enumerate(colunas, 1):
            folha.column_dimensions[get_column_letter(n)].width = largura
        folha.append([])
        for n, texto in enumerate(contexto_linhas(contexto), 2):
            celula = WriteOnlyCell(folha, value=limpar(texto))
            celula.data_type = "s"
            celula.font = Font(name="Arial", size=14 if n == 2 else 10, bold=n == 2)
            folha.append([celula])
        cabecalhos = []
        for _, titulo, _ in colunas:
            celula = WriteOnlyCell(folha, value=titulo)
            celula.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="1E3A5F")
            celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cabecalhos.append(celula)
        folha.row_dimensions[7].height = 32
        folha.append(cabecalhos)
        for numero, registro in enumerate(registros, 8):
            celulas, linhas = [], 1
            for chave, _, largura in colunas:
                valor = valor_celula(registro, chave)
                celula = WriteOnlyCell(folha, value=valor)
                if isinstance(valor, str):
                    celula.data_type = "s"
                    linhas = max(linhas, sum(max(1, ceil(len(p) / largura)) for p in valor.split("\n")))
                celula.font = Font(name="Arial", size=10, bold=registro.get("unidade") == "TOTAL")
                celula.alignment = Alignment(vertical="top", wrap_text=True, horizontal="right" if isinstance(valor, (int, float)) else "left")
                if numero % 2 == 0:
                    celula.fill = PatternFill("solid", fgColor="F1F5F9")
                if isinstance(valor, datetime):
                    celula.number_format = "dd/mm/yyyy hh:mm:ss"
                celulas.append(celula)
            folha.row_dimensions[numero].height = min(409, max(24, 14 * linhas + 8))
            folha.append(celulas)
    buffer = BytesIO()
    livro.save(buffer)
    return buffer.getvalue()


def ods(itens, resumo, totais, contexto):
    livro = OpenDocumentSpreadsheet()
    formato_data = DateStyle(name="DataHoraBrasilia")
    for elemento in (Day(style="long"), Text(text="/"), Month(style="long"), Text(text="/"), Year(style="long"),
        Text(text=" "), Hours(style="long"), Text(text=":"), Minutes(style="long"), Text(text=":"), Seconds(style="long")):
        formato_data.addElement(elemento)
    livro.styles.addElement(formato_data)
    for nome, fundo, negrito in (("Normal", None, False), ("Alternada", "#f1f5f9", False), ("Cabecalho", "#1e3a5f", True)):
        estilo = Style(name=nome, family="table-cell")
        estilo.addElement(TextProperties(fontfamily="Arial", fontsize="10pt", fontweight="bold" if negrito else "normal", color="#ffffff" if negrito else "#0f172a"))
        estilo.addElement(TableCellProperties(wrapoption="wrap", verticalalign="top", **({"backgroundcolor": fundo} if fundo else {})))
        livro.styles.addElement(estilo)
        if nome != "Cabecalho":
            data_estilo = Style(name=nome + "DataHora", family="table-cell", parentstylename=estilo, datastylename=formato_data)
            livro.styles.addElement(data_estilo)
    for nome, tamanho in (("Contexto", "10pt"), ("Titulo", "14pt")):
        estilo = Style(name=nome, family="table-cell")
        estilo.addElement(TextProperties(fontfamily="Arial", fontsize=tamanho, fontweight="bold" if nome == "Titulo" else "normal"))
        estilo.addElement(TableCellProperties(wrapoption="no-wrap"))
        livro.styles.addElement(estilo)
    altura = Style(name="AlturaAutomatica", family="table-row")
    altura.addElement(TableRowProperties(useoptimalrowheight="true", minrowheight="0.55cm"))
    livro.styles.addElement(altura)
    for nome, colunas, registros in tabelas(itens, resumo, totais):
        tabela = Table(name=nome)
        for n, (_, _, largura) in enumerate(colunas):
            estilo = Style(name=f"Coluna{nome}{n}", family="table-column")
            estilo.addElement(TableColumnProperties(columnwidth=f"{largura * 0.19:.2f}cm"))
            livro.automaticstyles.addElement(estilo)
            tabela.addElement(TableColumn(stylename=estilo))

        def adicionar(valores, estilo="Normal"):
            linha = TableRow(stylename=altura)
            for valor in valores:
                if isinstance(valor, datetime):
                    celula = TableCell(valuetype="date", datevalue=valor.isoformat(), stylename=estilo + "DataHora")
                    texto = valor.strftime("%d/%m/%Y %H:%M:%S")
                elif isinstance(valor, (int, float)):
                    celula = TableCell(valuetype="float", value=valor, stylename=estilo)
                    texto = str(valor)
                else:
                    celula = TableCell(valuetype="string", stylename=estilo)
                    texto = limpar(str(valor)) if valor is not None else ""
                for parte in texto.split("\n"):
                    celula.addElement(P(text=parte))
                linha.addElement(celula)
            tabela.addElement(linha)

        adicionar([])
        for n, linha in enumerate(contexto_linhas(contexto)):
            adicionar([linha], "Titulo" if n == 0 else "Contexto")
        adicionar([titulo for _, titulo, _ in colunas], "Cabecalho")
        for n, registro in enumerate(registros):
            adicionar([valor_celula(registro, chave) for chave, _, _ in colunas], "Alternada" if n % 2 == 0 else "Normal")
        livro.spreadsheet.addElement(tabela)
    buffer = BytesIO()
    livro.write(buffer)
    return buffer.getvalue()


def csv_sheets(itens):
    buffer = StringIO(newline="")
    escritor = csv.writer(buffer, delimiter=",")
    escritor.writerow([titulo for _, titulo, _ in COLUNAS])
    for item in itens:
        linha = []
        for chave, _, _ in COLUNAS:
            valor = valor_celula(item, chave)
            if isinstance(valor, datetime):
                valor = valor.isoformat(sep=" ")
            if isinstance(valor, str) and valor.lstrip().startswith(("=", "+", "-", "@")):
                valor = "'" + valor
            linha.append(valor)
        escritor.writerow(linha)
    return buffer.getvalue().encode("utf-8-sig")


def exportar_dietas(formato, itens, resumo, totais, contexto):
    if formato == "xlsx" and any(isinstance(valor_celula(item, chave), str) and len(valor_celula(item, chave)) > 32767
        for item in itens for chave, _, _ in COLUNAS):
        raise ValueError("Um campo excede o limite de texto do Excel. Exporte em CSV ou ODS para preservar o conteúdo completo")
    gerador = {"xlsx": xlsx, "ods": ods}.get(formato)
    arquivo = gerador(itens, resumo, totais, contexto) if gerador else csv_sheets(itens)
    return arquivo, MIMES[formato]
